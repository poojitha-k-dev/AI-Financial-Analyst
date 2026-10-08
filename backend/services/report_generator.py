"""
PDF & Excel report generator.
"""
import io
import os
from datetime import datetime
from typing import Dict, Any, Optional

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    HRFlowable, PageBreak,
)
from reportlab.lib.enums import TA_CENTER, TA_LEFT

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_pdf_report(
    company_name: str,
    period: str,
    financial_data: Dict[str, Any],
    kpis: Dict[str, Any],
    risk_data: Dict[str, Any],
    insights_text: str,
    output_path: str,
) -> str:
    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=1 * inch,
        bottomMargin=0.75 * inch,
    )

    styles = getSampleStyleSheet()
    story = []

    # ── Custom styles ──
    title_style = ParagraphStyle(
        "TitleStyle",
        parent=styles["Title"],
        fontSize=22,
        textColor=colors.HexColor("#1a1a2e"),
        spaceAfter=6,
    )
    h1_style = ParagraphStyle(
        "H1Style",
        parent=styles["Heading1"],
        fontSize=14,
        textColor=colors.HexColor("#16213e"),
        spaceBefore=16,
        spaceAfter=6,
        borderPad=4,
    )
    h2_style = ParagraphStyle(
        "H2Style",
        parent=styles["Heading2"],
        fontSize=11,
        textColor=colors.HexColor("#0f3460"),
        spaceBefore=10,
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "BodyStyle",
        parent=styles["Normal"],
        fontSize=9,
        leading=14,
        spaceAfter=4,
    )
    small_style = ParagraphStyle(
        "SmallStyle",
        parent=styles["Normal"],
        fontSize=8,
        textColor=colors.grey,
    )

    # ── Cover ──
    story.append(Spacer(1, 0.5 * inch))
    story.append(Paragraph("AI FINANCIAL ANALYST", ParagraphStyle(
        "Brand", parent=styles["Normal"], fontSize=10,
        textColor=colors.HexColor("#e94560"), alignment=TA_CENTER
    )))
    story.append(Spacer(1, 0.2 * inch))
    story.append(Paragraph(f"Financial Analysis Report", title_style))
    story.append(Paragraph(f"{company_name}", ParagraphStyle(
        "Co", parent=styles["Normal"], fontSize=16,
        textColor=colors.HexColor("#0f3460"), alignment=TA_LEFT
    )))
    story.append(Paragraph(f"Period: {period}", body_style))
    story.append(Paragraph(f"Generated: {datetime.now().strftime('%B %d, %Y %H:%M')}", small_style))
    story.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor("#e94560")))
    story.append(Spacer(1, 0.3 * inch))

    # ── Key Metrics Summary Table ──
    story.append(Paragraph("Key Financial Metrics", h1_style))
    metrics_rows = [["Metric", "Value", "Category"]]
    flat_metrics = [
        ("Revenue", financial_data.get("revenue"), "Income"),
        ("Net Income", financial_data.get("net_income"), "Income"),
        ("EBITDA", financial_data.get("ebitda"), "Income"),
        ("Total Assets", financial_data.get("total_assets"), "Balance Sheet"),
        ("Total Equity", financial_data.get("total_equity"), "Balance Sheet"),
        ("Total Debt", financial_data.get("total_debt"), "Balance Sheet"),
        ("Operating Cash Flow", financial_data.get("operating_cash_flow"), "Cash Flow"),
        ("Free Cash Flow", financial_data.get("free_cash_flow"), "Cash Flow"),
    ]
    for name, val, cat in flat_metrics:
        if val is not None:
            metrics_rows.append([name, _fmt_currency(val), cat])

    t = Table(metrics_rows, colWidths=[2.5 * inch, 2 * inch, 1.5 * inch])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16213e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8f9fa")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#dee2e6")),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("PADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.3 * inch))

    # ── KPI Tables ──
    story.append(Paragraph("KPI Analysis", h1_style))
    kpi_categories = {
        "Profitability": kpis.get("profitability", {}),
        "Liquidity": kpis.get("liquidity", {}),
        "Leverage": kpis.get("leverage", {}),
        "Efficiency": kpis.get("efficiency", {}),
        "Cash Flow": kpis.get("cash_flow", {}),
    }

    for cat_name, cat_kpis in kpi_categories.items():
        if not cat_kpis:
            continue
        story.append(Paragraph(cat_name, h2_style))
        rows = [["KPI", "Value"]]
        for k, v in cat_kpis.items():
            if v is not None:
                label = k.replace("_", " ").title()
                suffix = "%" if any(x in k for x in ["margin", "growth", "rate", "return"]) else "x" if any(x in k for x in ["ratio", "turnover", "coverage"]) else ""
                rows.append([label, f"{v:.2f}{suffix}"])
        if len(rows) > 1:
            t2 = Table(rows, colWidths=[3 * inch, 2 * inch])
            t2.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f3460")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f4f8")]),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#dee2e6")),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("PADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(t2)
            story.append(Spacer(1, 0.15 * inch))

    story.append(PageBreak())

    # ── Risk Section ──
    story.append(Paragraph("Risk Analysis", h1_style))
    comp_scores = kpis.get("composite_scores", {})
    if comp_scores:
        risk_rows = [["Risk Dimension", "Score", "Rating"]]
        for k, v in comp_scores.items():
            if v is not None:
                label = k.replace("_", " ").title()
                rating = "Strong" if v >= 7 else "Moderate" if v >= 4 else "Weak"
                risk_rows.append([label, f"{v:.1f}/10", rating])
        t3 = Table(risk_rows, colWidths=[2.5 * inch, 1.5 * inch, 1.5 * inch])
        t3.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e94560")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fff5f5")]),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#dee2e6")),
            ("PADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(t3)
        story.append(Spacer(1, 0.2 * inch))

    if risk_data:
        altman = risk_data.get("altman_z_score", {})
        if altman.get("score"):
            story.append(Paragraph(f"Altman Z-Score: {altman['score']:.2f} – {altman.get('zone', '')}", h2_style))
            story.append(Paragraph(altman.get("interpretation", ""), body_style))

    story.append(PageBreak())

    # ── AI Insights ──
    story.append(Paragraph("AI-Generated Financial Insights", h1_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#e94560")))
    story.append(Spacer(1, 0.1 * inch))

    # Parse markdown-ish text into paragraphs
    for line in insights_text.split("\n"):
        line = line.strip()
        if not line:
            story.append(Spacer(1, 0.05 * inch))
        elif line.startswith("## "):
            story.append(Paragraph(line[3:], h1_style))
        elif line.startswith("### "):
            story.append(Paragraph(line[4:], h2_style))
        elif line.startswith("**") and line.endswith("**"):
            story.append(Paragraph(f"<b>{line[2:-2]}</b>", body_style))
        elif line.startswith("- ") or line.startswith("• "):
            story.append(Paragraph(f"• {line[2:]}", body_style))
        else:
            story.append(Paragraph(line, body_style))

    # ── Footer note ──
    story.append(Spacer(1, 0.5 * inch))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.grey))
    story.append(Paragraph(
        "This report was generated by AI Financial Analyst. "
        "It is for informational purposes only and does not constitute financial advice.",
        small_style,
    ))

    doc.build(story)
    return output_path


def generate_excel_report(
    company_name: str,
    period: str,
    financial_data: Dict[str, Any],
    kpis: Dict[str, Any],
    output_path: str,
) -> str:
    wb = openpyxl.Workbook()

    # ── Summary Sheet ──
    ws = wb.active
    ws.title = "Summary"
    _style_sheet_header(ws, f"{company_name} – {period}")

    row = 3
    ws.cell(row, 1, "FINANCIAL OVERVIEW").font = Font(bold=True, size=12, color="1a1a2e")
    row += 1

    headers = ["Metric", "Value"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row, col, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="16213e")
        c.alignment = Alignment(horizontal="center")
    row += 1

    metrics = [
        ("Revenue", financial_data.get("revenue")),
        ("Gross Profit", financial_data.get("gross_profit")),
        ("Operating Income", financial_data.get("operating_income")),
        ("Net Income", financial_data.get("net_income")),
        ("EBITDA", financial_data.get("ebitda")),
        ("Total Assets", financial_data.get("total_assets")),
        ("Total Equity", financial_data.get("total_equity")),
        ("Total Debt", financial_data.get("total_debt")),
        ("Operating Cash Flow", financial_data.get("operating_cash_flow")),
        ("Free Cash Flow", financial_data.get("free_cash_flow")),
    ]
    fill_alt = PatternFill("solid", fgColor="f8f9fa")
    for i, (name, val) in enumerate(metrics):
        if val is not None:
            ws.cell(row, 1, name)
            c = ws.cell(row, 2, val)
            c.number_format = '#,##0.00'
            if i % 2 == 0:
                for col in range(1, 3):
                    ws.cell(row, col).fill = fill_alt
            row += 1

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18

    # ── KPI Sheets ──
    kpi_map = {
        "Profitability": kpis.get("profitability", {}),
        "Liquidity": kpis.get("liquidity", {}),
        "Leverage": kpis.get("leverage", {}),
        "Efficiency": kpis.get("efficiency", {}),
        "Cash Flow": kpis.get("cash_flow", {}),
        "Growth": kpis.get("growth", {}),
    }

    for sheet_name, kpi_data in kpi_map.items():
        if not kpi_data:
            continue
        ws2 = wb.create_sheet(title=sheet_name)
        _style_sheet_header(ws2, f"{sheet_name} KPIs – {company_name}")

        row = 3
        ws2.cell(row, 1, "KPI").font = Font(bold=True, color="FFFFFF")
        ws2.cell(row, 1).fill = PatternFill("solid", fgColor="0f3460")
        ws2.cell(row, 2, "Value").font = Font(bold=True, color="FFFFFF")
        ws2.cell(row, 2).fill = PatternFill("solid", fgColor="0f3460")
        row += 1

        for i, (k, v) in enumerate(kpi_data.items()):
            if v is not None:
                ws2.cell(row, 1, k.replace("_", " ").title())
                c = ws2.cell(row, 2, v)
                c.number_format = "0.00"
                if i % 2 == 0:
                    for col in range(1, 3):
                        ws2.cell(row, col).fill = PatternFill("solid", fgColor="f0f4f8")
                row += 1

        ws2.column_dimensions["A"].width = 30
        ws2.column_dimensions["B"].width = 15

    wb.save(output_path)
    return output_path


def _style_sheet_header(ws, title: str):
    ws.merge_cells("A1:D1")
    title_cell = ws.cell(1, 1, title)
    title_cell.font = Font(bold=True, size=14, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="1a1a2e")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28


def _fmt_currency(val: float) -> str:
    if abs(val) >= 1e9:
        return f"${val/1e9:.2f}B"
    elif abs(val) >= 1e6:
        return f"${val/1e6:.2f}M"
    elif abs(val) >= 1e3:
        return f"${val/1e3:.2f}K"
    return f"${val:.2f}"
